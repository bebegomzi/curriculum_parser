"""
교육과정 편성표(.xlsx)를 LLM이 읽기 쉬운 CSV 3종으로 변환합니다.

전제
- 헤더/열 배열은 자동으로 감지되거나 기본 구조를 유지하며, 과목/요약 테이블의 길이는 달라도 동작합니다.
- 병합 셀은 좌상단 값을 병합 범위 전체에 전파합니다.
- 각 학기 칸의 숫자는 학점/주당 시수/이수 단위로 해석합니다.

출력
1) *_courses_llm.csv           : 과목별 평면표
2) *_selection_groups_llm.csv  : 학생 선택군 요약
3) *_summary_llm.csv           : 하단 합계/창체 요약
4) *_readme.md                 : 해석 규칙

사용 예
python convert_curriculum_to_llm_csv.py "2026학년도 반곡고등학교 입학생 3개년 교육과정 편성표.xlsx"
"""

from __future__ import annotations

import argparse
import csv
import re
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


# ===== 고정 레이아웃 설정 (기본값) =====

COL = {
    "교과군": 1,                 # A
    "과목명": 2,                 # B
    "학교지정": 3,               # C
    "학년선택": 4,               # D
    "공통": 5,                   # E
    "일반": 6,                   # F
    "진로": 7,                   # G
    "융합": 8,                   # H
    "비고_전문_고시외_표6": 9,   # I
    "학생선택_1학기": 10,        # J
    "학생선택_2학기": 11,        # K
    "기본학점": 12,              # L
    "1-1": 13,                   # M
    "1-2": 14,                   # N
    "2-1": 15,                   # O
    "2-2": 16,                   # P
    "3-1": 17,                   # Q
    "3-2": 18,                   # R
    "편성학점": 19,              # S
    "교과군별이수학점": 20,      # T
    "필수이수학점": 21,          # U
    "과학중점": 22,              # V
    "정보중점": 23,              # W
    "사회중점": 24,              # X
    "인문중점": 25,              # Y
    "캠공개설여부": 26,          # Z
}

SEMESTER_COLS = {
    "1-1": 13,
    "1-2": 14,
    "2-1": 15,
    "2-2": 16,
    "3-1": 17,
    "3-2": 18,
}

COURSE_FIELDNAMES = [
    "원본행",
    "교과군",
    "과목명",
    "학교지정",
    "학년선택",
    "과목유형",
    "비고_전문_고시외_표6",
    "학생선택_원문",
    "선택군_코드",
    "선택군_학년학기",
    "선택군_선택수",
    "선택군_과목당학점",
    "선택군_총학점",
    "선택군_행범위",
    "기본학점",
    "1-1_원본",
    "1-2_원본",
    "2-1_원본",
    "2-2_원본",
    "3-1_원본",
    "3-2_원본",
    "1-1_학점",
    "1-2_학점",
    "2-1_학점",
    "2-2_학점",
    "3-1_학점",
    "3-2_학점",
    "편성학점",
    "교과군별이수학점",
    "필수이수학점",
    "과학중점",
    "정보중점",
    "사회중점",
    "인문중점",
    "캠공개설여부",
    "비고",
]

SELECTION_FIELDNAMES = [
    "선택군_코드",
    "선택군_학년학기",
    "선택군_원문",
    "선택군_선택수",
    "선택군_과목당학점",
    "선택군_총학점",
    "선택군_행범위",
    "대상과목수",
    "대상과목목록",
]

SUMMARY_FIELDNAMES = [
    "원본행",
    "항목",
    "세부항목",
    "1-1",
    "1-2",
    "2-1",
    "2-2",
    "3-1",
    "3-2",
    "편성",
    "교과군별이수학점",
    "필수이수학점",
    "비고",
]


@dataclass
class SelectionGroup:
    code: str
    semester: str
    raw: str
    choose_count: str
    credit_per_course: str
    total_credit: str
    row_start: int
    row_end: int
    subjects: List[str]

    @property
    def row_range(self) -> str:
        return f"{self.row_start}:{self.row_end}"

    def as_row(self) -> Dict[str, str]:
        return {
            "선택군_코드": self.code,
            "선택군_학년학기": self.semester,
            "선택군_원문": self.raw,
            "선택군_선택수": self.choose_count,
            "선택군_과목당학점": self.credit_per_course,
            "선택군_총학점": self.total_credit,
            "선택군_행범위": self.row_range,
            "대상과목수": str(len(self.subjects)),
            "대상과목목록": "; ".join(self.subjects),
        }


# ===== 한글 로마자 표기 헬퍼 (접두사 자동 생성용) =====

CHOSEONG = [
    'g', 'kk', 'n', 'd', 'tt', 'r', 'm', 'b', 'pp',
    's', 'ss', '', 'j', 'jj', 'ch', 'k', 't', 'p', 'h'
]
JUNGSEONG = [
    'a', 'ae', 'ya', 'yae', 'eo', 'e', 'yeo', 'ye', 'o',
    'wa', 'wae', 'oe', 'yo', 'u', 'wo', 'we', 'wi', 'yu',
    'eu', 'ui', 'i'
]
JONGSEONG = [
    '', 'k', 'k', 'k', 'n', 'n', 'n', 't', 'l', 'lg',
    'lm', 'lp', 'lh', 'lh', 'lp', 'lh', 'm', 'p', 'p',
    't', 't', 'ng', 't', 't', 'k', 't', 'p', 't'
]

def romanize_char(char: str) -> str:
    code = ord(char)
    if 0xAC00 <= code <= 0xD7A3:
        temp = code - 0xAC00
        cho = temp // 588
        jung = (temp % 588) // 28
        jong = temp % 28
        return CHOSEONG[cho] + JUNGSEONG[jung] + JONGSEONG[jong]
    return char.lower()

def romanize(text: str) -> str:
    res = []
    for char in text:
        if 0xAC00 <= ord(char) <= 0xD7A3:
            res.append(romanize_char(char))
        elif char.isalnum():
            res.append(char.lower())
        elif char in (' ', '_', '-'):
            res.append('_')
    out = "".join(res)
    out = re.sub(r'_+', '_', out)
    return out.strip('_')


# ===== 공통 유틸리티 함수 =====

def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_space(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def cell_val(ws: Worksheet, row: int, col_name: str) -> Any:
    col_idx = COL.get(col_name)
    if col_idx is None or col_idx < 1:
        return None
    return ws.cell(row, col_idx).value


def true_marker(value: Any) -> str:
    """○ 또는 이미 True로 볼 수 있는 값만 True, 나머지는 빈칸."""
    text = clean(value)
    return "True" if text in {"○", "TRUE", "True", "true", "1"} else ""


def number_text(value: Any) -> str:
    """CSV에 넣을 숫자 텍스트. 비숫자 메모는 그대로 보존."""
    return clean(value)


def extract_first_number(value: Any) -> str:
    """학기 칸에서 첫 숫자만 학점으로 추출. 예: '3\n(공동)' -> '3'."""
    text = clean(value)
    if not text:
        return ""
    m = re.search(r"\d+(?:\.\d+)?", text)
    if not m:
        return ""
    n = float(m.group())
    return str(int(n)) if n.is_integer() else str(n)


def copy_ws_with_filled_merges(ws: Worksheet) -> Worksheet:
    """
    원본 시트를 메모리상에서 복사한 뒤 병합 셀을 해제하고 값을 전파합니다.
    원본 파일은 변경하지 않습니다.
    """
    new_ws = copy(ws)
    for merged in list(new_ws.merged_cells.ranges):
        value = new_ws.cell(merged.min_row, merged.min_col).value
        new_ws.unmerge_cells(str(merged))
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                new_ws.cell(row, col).value = value
    return new_ws


# ===== 동적 탐색 및 감지 로직 =====

def detect_curriculum_sheet(wb: load_workbook) -> str:
    """학업 교육과정 편성표의 특징적 헤더가 포함된 시트를 찾아 반환합니다."""
    # Prioritize sheet names with keywords
    priority_sheets = []
    for name in wb.sheetnames:
        name_clean = name.lower()
        score = 0
        if "변경후" in name_clean:
            score += 10
        if "입학생" in name_clean:
            score += 5
        if "1학년" in name_clean:
            score += 3
        if score > 0:
            priority_sheets.append((score, name))
    
    priority_sheets.sort(key=lambda x: x[0], reverse=True)
    search_order = [name for _, name in priority_sheets] + [n for n in wb.sheetnames if n not in [s[1] for s in priority_sheets]]

    for name in search_order:
        ws = wb[name]
        if ws.max_row >= 8 and ws.max_column >= 12:
            # 첫 15행을 뒤져서 헤더 행을 탐색해 봅니다
            for r in range(1, 16):
                row_vals = [clean(ws.cell(r, c).value) for c in range(1, min(ws.max_column + 1, 25))]
                has_gyogwa = any("교과" in val for val in row_vals)
                has_gwamok = any("과목" in val for val in row_vals)
                if has_gyogwa and has_gwamok:
                    return name
    return wb.sheetnames[0]


def find_header_start_row(ws: Worksheet) -> int:
    """헤더 행('교과(군)'과 '과목명/과목'이 나타나는 행)을 찾아 반환합니다."""
    for r in range(1, 21):
        row_vals = [clean(ws.cell(r, c).value) for c in range(1, min(ws.max_column + 1, 25))]
        has_gyogwa = any("교과" in val for val in row_vals)
        has_gwamok = any("과목" in val for val in row_vals)
        if has_gyogwa and has_gwamok:
            return r
    return 5  # 기본 fallback


def detect_columns(ws: Worksheet, header_start: int) -> Dict[str, int]:
    """3줄의 헤더 행 텍스트 조합을 스캔하여 각각의 데이터 열 인덱스를 동적으로 매핑합니다."""
    global SEMESTER_COLS
    col_map = COL.copy()
    
    # Check if Wongsang style by reading the first cell in header_start row
    first_cell = clean(ws.cell(header_start, 1).value)
    if first_cell == "구분":
        # Wongsang layout is fixed!
        col_map["구분"] = 1
        col_map["교과군"] = 2
        col_map["과목유형_단일열"] = 3
        col_map["과목명"] = 4
        col_map["기본학점"] = 5
        col_map["편성학점"] = 6
        col_map["1-1"] = 7
        col_map["1-2"] = 8
        col_map["2-1"] = 9
        col_map["2-2"] = 10
        col_map["3-1"] = 11
        col_map["3-2"] = 12
        col_map["편성학점합"] = 13
        col_map["교과군별이수학점"] = 13  # fallback
        
        # Update SEMESTER_COLS
        SEMESTER_COLS = {
            "1-1": 7,
            "1-2": 8,
            "2-1": 9,
            "2-2": 10,
            "3-1": 11,
            "3-2": 12,
        }
        return col_map

    col_headers = {}
    for c in range(1, ws.max_column + 1):
        vals = []
        for r in range(header_start, header_start + 3):
            val = ws.cell(r, c).value
            if val is not None:
                vals.append(str(val).strip())
        col_headers[c] = " ".join(vals)

    # 매핑 규칙 정의
    rules = {
        "구분": lambda text: text.strip() == "구분",
        "교과군": lambda text: "교과" in text,
        "과목명": lambda text: "과목" in text and "유형" not in text and "코드" not in text,
        "학교지정": lambda text: "구분" in text and "지정" in text,
        "학년선택": lambda text: "구분" in text and "선택" in text,
        "과목유형_단일열": lambda text: "과목" in text and "유형" in text,
        "공통": lambda text: "공통" in text and "과목유형" in text,
        "일반": lambda text: "일반" in text and "과목유형" in text,
        "진로": lambda text: "진로" in text and "과목유형" in text,
        "융합": lambda text: "융합" in text and "과목유형" in text,
        "비고_전문_고시외_표6": lambda text: any(x in text for x in ["전문", "고시외", "표6"]),
        "학생선택_1학기": lambda text: "학생선택" in text and text.endswith("1학기"),
        "학생선택_2학기": lambda text: "학생선택" in text and text.endswith("2학기"),
        "기본학점": lambda text: "기본학점" in text,
        "1-1": lambda text: ("1학년" in text and "1학기" in text) or "1-1" in text or "1학년1학기" in text or "1학년 1학기" in text,
        "1-2": lambda text: ("1학년" in text and "2학기" in text) or "1-2" in text or "1학년2학기" in text or "1학년 2학기" in text,
        "2-1": lambda text: ("2학년" in text and "1학기" in text) or "2-1" in text or "2학년1학기" in text or "2학년 1학기" in text,
        "2-2": lambda text: ("2학년" in text and "2학기" in text) or "2-2" in text or "2학년2학기" in text or "2학년 2학기" in text,
        "3-1": lambda text: ("3학년" in text and "1학기" in text) or "3-1" in text or "3학년1학기" in text or "3학년 1학기" in text,
        "3-2": lambda text: ("3학년" in text and "2학기" in text) or "3-2" in text or "3학년2학기" in text or "3학년 2학기" in text,
        "편성학점": lambda text: ("편성" in text or "운영" in text) and "합" not in text and "이수" not in text and "필수" not in text,
        "편성학점합": lambda text: "편성" in text and "합" in text,
        "교과군별이수학점": lambda text: "교과군별" in text or "교과군별이수" in text,
        "필수이수학점": lambda text: "필수" in text and "이수" in text,
        "과학중점": lambda text: "중점" in text and text.endswith("과학"),
        "정보중점": lambda text: "중점" in text and text.endswith("정보"),
        "사회중점": lambda text: "중점" in text and text.endswith("사회"),
        "인문중점": lambda text: "중점" in text and text.endswith("인문"),
        "캠공개설여부": lambda text: "캠공" in text or "공동교육" in text or "개설여부" in text,
    }

    # 스캔하여 규칙 매칭
    detected = {}
    for key, rule in rules.items():
        for c, text in col_headers.items():
            if rule(text):
                detected[key] = c
                break

    # fallback for 교과군별이수학점
    if "교과군별이수학점" not in detected and "편성학점합" in detected:
        detected["교과군별이수학점"] = detected["편성학점합"]

    col_map.update(detected)

    # Clean up single-column layout if checkboxes are present
    if any(k in detected for k in ["공통", "일반", "진로", "융합"]):
        if "과목유형_단일열" in col_map:
            del col_map["과목유형_단일열"]

    # 학기별 열 정보 업데이트
    SEMESTER_COLS = {
        "1-1": col_map.get("1-1", 13),
        "1-2": col_map.get("1-2", 14),
        "2-1": col_map.get("2-1", 15),
        "2-2": col_map.get("2-2", 16),
        "3-1": col_map.get("3-1", 17),
        "3-2": col_map.get("3-2", 18),
    }

    return col_map


def extract_school_and_year(input_path: Path, ws: Worksheet) -> Tuple[Optional[str], Optional[str]]:
    """파일명과 시트 내용에서 학교명(예: 반곡)과 입학년도(예: 2026)를 추출합니다."""
    year = None
    school = None

    # 1. 파일명에서 탐색
    filename = input_path.name
    m_year = re.search(r"(\d{4})학년도", filename)
    if m_year:
        year = m_year.group(1)
    
    m_school = re.search(r"([가-힣\w]+(?:고등학교|고))", filename)
    if m_school:
        school = m_school.group(1)

    # 2. 파일명에 없으면 시트 본문 상단 5행을 뒤져서 탐색
    if not year or not school:
        for r in range(1, 6):
            for col in range(1, ws.max_column + 1):
                val = clean(ws.cell(r, col).value)
                if not val:
                    continue
                if not year:
                    m_y = re.search(r"(\d{4})학년도", val)
                    if m_y:
                        year = m_y.group(1)
                if not school:
                    m_s = re.search(r"([가-힣\w]+(?:고등학교|고))", val)
                    if m_s:
                        school = m_s.group(1)

    # 3. 그래도 학교명이 없으면 시트명에서 추출
    if not school:
        m_s = re.search(r"([가-힣\w]+(?:고등학교|고))", ws.title)
        if m_s:
            school = m_s.group(1)
        else:
            school = ws.title

    # 4. 고등학교/고 접두사 클리닝
    if school:
        school_clean = re.sub(r"(고등학교|고)$", "", school)
        if school_clean:
            school = school_clean

    return school, year


def update_column_mapping(ws: Worksheet, header_start: int) -> None:
    global COL
    detected = detect_columns(ws, header_start)
    COL.update(detected)


# ===== 비즈니스 파싱 로직 =====

def find_summary_start_row(ws: Worksheet, data_start: int) -> int:
    """
    과목 목록이 끝나고 하단 요약이 시작되는 행을 찾습니다.
    A:L 병합 합계행 또는 과목명 소멸 지점을 기준으로 잡습니다.
    """
    for r in range(data_start, ws.max_row + 1):
        a = clean(cell_val(ws, r, "교과군"))
        b = clean(cell_val(ws, r, "과목명"))
        joined = normalize_space(" ".join([a, b]))
        if any(keyword in joined for keyword in ["창의적", "총 이수", "총계", "합계", "학생 선택", "총 교과"]):
            return r
        # 과목명 없이 학기/편성 쪽 숫자만 나오기 시작하면 요약 영역으로 간주
        if r > data_start and not b:
            col_1_1 = COL.get("1-1", 13)
            col_required = COL.get("필수이수학점", 21)
            right_values = [clean(ws.cell(r, c).value) for c in range(col_1_1, col_required + 1)]
            if any(right_values):
                return r
    return ws.max_row + 1


def course_type(ws: Worksheet, row: int) -> str:
    """공통/일반/진로/융합 중 주 과목유형을 반환합니다."""
    col_single = COL.get("과목유형_단일열")
    if col_single and col_single >= 1:
        return clean(ws.cell(row, col_single).value)
    for name in ["공통", "일반", "진로", "융합"]:
        if true_marker(cell_val(ws, row, name)):
            return name
    return ""


def split_note_suffix(raw: Any) -> str:
    return clean(raw)


def parse_selection_meta(raw: str) -> Tuple[str, str, str]:
    text = clean(raw)
    choose = ""
    credit = ""

    m_choose = re.search(r"택\s*(\d+)", text)
    if m_choose:
        choose = m_choose.group(1)

    m_credit = re.search(r"\((\d+(?:\.\d+)?)\)", text)
    if m_credit:
        credit = m_credit.group(1)

    total = ""
    if choose and credit:
        total_num = float(choose) * float(credit)
        total = str(int(total_num)) if total_num.is_integer() else str(total_num)

    return choose, credit, total


def compact_selection_label(raw: str) -> str:
    first_line = clean(raw).splitlines()[0] if clean(raw) else ""
    return re.sub(r"\s+", "", first_line)


def detect_selection_group_for_row_wongsang(ws: Worksheet, row: int, data_start: int) -> Tuple[str, str, str, str, str, str, str]:
    col_gubun = COL.get("구분")
    if not col_gubun:
        return "", "", "", "", "", "", ""

    group_name = clean(ws.cell(row, col_gubun).value)
    if not group_name:
        return "", "", "", "", "", "", ""

    row_start = row
    row_end = row
    while row_start > data_start and clean(ws.cell(row_start - 1, col_gubun).value) == group_name:
        row_start -= 1
    while row_end < ws.max_row and clean(ws.cell(row_end + 1, col_gubun).value) == group_name:
        row_end += 1
    row_range = f"{row_start}:{row_end}"

    semester = ""
    credit_val = ""
    for r in range(row_start, row_end + 1):
        for sem, col_idx in SEMESTER_COLS.items():
            val = clean(ws.cell(r, col_idx).value)
            if val:
                num = extract_first_number(val)
                if num:
                    semester = sem
                    credit_val = num
                    break
        if semester:
            break

    if not semester:
        return "", "", "", "", "", "", row_range

    col_total = COL.get("편성학점합")
    total_credit = ""
    if col_total:
        for r in range(row_start, row_end + 1):
            val = clean(ws.cell(r, col_total).value)
            if val:
                total_credit = val
                break

    if not total_credit:
        total_credit = credit_val

    try:
        choose_count = str(int(float(total_credit) // float(credit_val)))
    except:
        choose_count = "1"

    raw = f"택{choose_count} ({credit_val})"
    code = f"{semester}-{group_name}"

    return raw, code, semester, choose_count, credit_val, total_credit, row_range


def detect_selection_group_for_row(ws: Worksheet, row: int, data_start: int) -> Tuple[str, str, str, str, str, str, str]:
    """
    해당 과목 행의 학생선택 그룹을 추정합니다.
    J열 = 1학기 선택군, K열 = 2학기 선택군이지만,
    같은 '택4 (3)'이라도 실제 운영 학년-학기에 따라 별도 코드로 분리합니다.
    """
    sem_values = {}
    for sem in SEMESTER_COLS.keys():
        sem_values[sem] = clean(cell_val(ws, row, sem))

    candidates: List[Tuple[str, str]] = []
    for sem in ["1-1", "2-1", "3-1"]:
        if sem_values.get(sem):
            candidates.append((sem, clean(cell_val(ws, row, "학생선택_1학기"))))
    for sem in ["1-2", "2-2", "3-2"]:
        if sem_values.get(sem):
            candidates.append((sem, clean(cell_val(ws, row, "학생선택_2학기"))))

    candidates = [(sem, raw) for sem, raw in candidates if raw]
    if not candidates:
        return "", "", "", "", "", "", ""

    semester, raw = candidates[0]
    choose, credit, total = parse_selection_meta(raw)
    label = compact_selection_label(raw)
    code = f"{semester}-{label}" if label else ""

    sel_col = COL.get("학생선택_1학기") if semester.endswith("-1") else COL.get("학생선택_2학기")
    if not sel_col:
        return "", "", "", "", "", "", ""

    row_start = row
    row_end = row
    while row_start > data_start and clean(ws.cell(row_start - 1, sel_col).value) == raw:
        row_start -= 1
    while row_end < ws.max_row and clean(ws.cell(row_end + 1, sel_col).value) == raw:
        row_end += 1
    row_range = f"{row_start}:{row_end}"

    return raw, code, semester, choose, credit, total, row_range


def build_selection_groups(course_rows: List[Dict[str, str]]) -> List[Dict[str, str]]:
    groups: Dict[str, SelectionGroup] = {}
    for row in course_rows:
        code = row.get("선택군_코드", "")
        if not code:
            continue
        start, end = row["선택군_행범위"].split(":")
        if code not in groups:
            groups[code] = SelectionGroup(
                code=code,
                semester=row["선택군_학년학기"],
                raw=row["학생선택_원문"],
                choose_count=row["선택군_선택수"],
                credit_per_course=row["선택군_과목당학점"],
                total_credit=row["선택군_총학점"],
                row_start=int(start),
                row_end=int(end),
                subjects=[],
            )
        groups[code].subjects.append(row["과목명"])
        groups[code].row_start = min(groups[code].row_start, int(start))
        groups[code].row_end = max(groups[code].row_end, int(end))

    return [groups[k].as_row() for k in sorted(groups.keys())]


def build_course_rows(ws: Worksheet, summary_start: int, data_start: int) -> List[Dict[str, str]]:
    rows: List[Dict[str, str]] = []
    col_gubun = COL.get("구분")

    for r in range(data_start, summary_start):
        subject = clean(cell_val(ws, r, "과목명"))
        if not subject:
            continue

        # Skip subtotal and summary rows
        if any(x in subject for x in ["소계", "총계", "합계", "창의적 체험", "이수학점"]):
            continue

        is_designated = False
        is_selection = False
        if col_gubun and col_gubun >= 1:
            gubun_val = clean(ws.cell(r, col_gubun).value)
            gubun_clean = re.sub(r"\s+", "", gubun_val)
            if "지정" in gubun_clean:
                is_designated = True
            elif "선택" in gubun_clean:
                is_selection = True
        else:
            is_designated = true_marker(cell_val(ws, r, "학교지정")) == "True"
            is_selection = true_marker(cell_val(ws, r, "학년선택")) == "True"

        has_selection_cols = "구분" not in COL

        if is_selection:
            if not has_selection_cols:
                sel_raw, sel_code, sel_sem, sel_choose, sel_credit, sel_total, sel_range = detect_selection_group_for_row_wongsang(ws, r, data_start)
            else:
                sel_raw, sel_code, sel_sem, sel_choose, sel_credit, sel_total, sel_range = detect_selection_group_for_row(ws, r, data_start)
        else:
            sel_raw, sel_code, sel_sem, sel_choose, sel_credit, sel_total, sel_range = "", "", "", "", "", "", ""

        row: Dict[str, str] = {
            "원본행": str(r),
            "교과군": clean(cell_val(ws, r, "교과군")),
            "과목명": subject,
            "학교지정": "True" if is_designated else "",
            "학년선택": "True" if is_selection else "",
            "과목유형": course_type(ws, r),
            "비고_전문_고시외_표6": split_note_suffix(cell_val(ws, r, "비고_전문_고시외_표6")),
            "학생선택_원문": sel_raw,
            "선택군_코드": sel_code,
            "선택군_학년학기": sel_sem,
            "선택군_선택수": sel_choose,
            "선택군_과목당학점": sel_credit,
            "선택군_총학점": sel_total,
            "선택군_행범위": sel_range,
            "기본학점": number_text(cell_val(ws, r, "기본학점")),
            "편성학점": number_text(cell_val(ws, r, "편성학점")),
            "교과군별이수학점": number_text(cell_val(ws, r, "교과군별이수학점")),
            "필수이수학점": number_text(cell_val(ws, r, "필수이수학점")),
            "과학중점": true_marker(cell_val(ws, r, "과학중점")),
            "정보중점": true_marker(cell_val(ws, r, "정보중점")),
            "사회중점": true_marker(cell_val(ws, r, "사회중점")),
            "인문중점": true_marker(cell_val(ws, r, "인문중점")),
            "캠공개설여부": true_marker(cell_val(ws, r, "캠공개설여부")),
            "비고": "",
        }

        for sem in SEMESTER_COLS.keys():
            original = number_text(cell_val(ws, r, sem))
            row[f"{sem}_원본"] = original
            if is_selection and not has_selection_cols and sel_sem == sem:
                row[f"{sem}_학점"] = sel_credit
            else:
                row[f"{sem}_학점"] = extract_first_number(original)

        rows.append(row)

    return rows


def build_summary_rows(ws: Worksheet, summary_start: int) -> List[Dict[str, str]]:
    rows: List[Dict[str, str]] = []
    col_limit = COL.get("필수이수학점")
    if not col_limit:
        col_limit = COL.get("편성학점합")
    if not col_limit:
        col_limit = 21

    col_gubun = COL.get("구분")

    for r in range(summary_start, ws.max_row + 1):
        values = [clean(ws.cell(r, c).value) for c in range(1, col_limit + 1)]
        if not any(values):
            continue

        hangmok_val = clean(cell_val(ws, r, "교과군"))
        if not hangmok_val and col_gubun:
            hangmok_val = clean(ws.cell(r, col_gubun).value)

        row = {
            "원본행": str(r),
            "항목": hangmok_val,
            "세부항목": clean(cell_val(ws, r, "과목명")),
            "1-1": number_text(cell_val(ws, r, "1-1")),
            "1-2": number_text(cell_val(ws, r, "1-2")),
            "2-1": number_text(cell_val(ws, r, "2-1")),
            "2-2": number_text(cell_val(ws, r, "2-2")),
            "3-1": number_text(cell_val(ws, r, "3-1")),
            "3-2": number_text(cell_val(ws, r, "3-2")),
            "편성": number_text(cell_val(ws, r, "편성학점")),
            "교과군별이수학점": number_text(cell_val(ws, r, "교과군별이수학점")),
            "필수이수학점": number_text(cell_val(ws, r, "필수이수학점")),
            "비고": "",
        }
        if any(v for k, v in row.items() if k != "원본행"):
            rows.append(row)
    return rows


def write_csv(path: Path, fieldnames: List[str], rows: Iterable[Dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def write_readme(path: Path, source_file: str, sheet: str, course_count: int, group_count: int, summary_count: int) -> None:
    text = f"""# 교육과정 편성표 LLM용 CSV 해석 규칙

## 원본
- 파일: `{source_file}`
- 시트: `{sheet}`

## 생성 파일
- `*_courses_llm.csv`: 과목 1개를 1행으로 정리한 평면표
- `*_selection_groups_llm.csv`: 학생 선택군 단위 요약표
- `*_summary_llm.csv`: 하단 합계/창의적 체험활동 등 요약 영역
- `*_readme.md`: 해석 규칙 설명 (본 파일)

## 변환 원칙
1. 병합 셀은 좌상단 값을 병합 범위 전체에 전파했다.
2. `1-1`, `1-2`, `2-1`, `2-2`, `3-1`, `3-2`의 숫자는 해당 학기 편성 학점이며, 이 양식에서는 주당 시수/이수 단위와 같은 의미로 해석한다.
3. `○` 표시는 `True`로 변환하고, 해당하지 않는 칸은 빈칸으로 둔다.
4. `과목유형`은 `공통`, `일반`, `진로`, `융합` 중 하나로 정리한다.
5. `전문`, `고시외`, `표6` 등은 과목유형이 아니라 `비고_전문_고시외_표6`에 별도로 둔다.
6. 학생 선택군은 같은 원문이라도 운영 학년-학기가 다르면 별도 선택군으로 본다. 예: `2-1-택4(3)`과 `2-2-택4(3)`은 다른 선택군이다.
7. `과학중점`, `정보중점`, `사회중점`, `인문중점`은 별도 boolean 열로 분리한다.

## 생성 결과 개수
- 과목 행: {course_count}
- 선택군: {group_count}
- 요약 행: {summary_count}
"""
    path.write_text(text, encoding="utf-8")


def convert(input_path: Path, sheet_name: Optional[str], outdir: Path, prefix: Optional[str]) -> Tuple[Dict[str, Path], str, str]:
    wb = load_workbook(input_path, data_only=True)
    
    # 1. 시트 자동 감지
    if not sheet_name:
        sheet_name = detect_curriculum_sheet(wb)
    elif sheet_name not in wb.sheetnames:
        raise ValueError(f"시트 '{sheet_name}'을 찾을 수 없습니다. 사용 가능 시트: {wb.sheetnames}")

    ws = copy_ws_with_filled_merges(wb[sheet_name])
    
    # 2. 헤더 시작 행 감지 및 열 매핑 업데이트
    header_start = find_header_start_row(ws)
    update_column_mapping(ws, header_start)
    
    # 3. 학교 및 연도 감지
    school, year = extract_school_and_year(input_path, ws)
    
    # 4. 파일 접두사 설정
    if not prefix:
        parts = []
        if school:
            parts.append(romanize(school))
        if year:
            parts.append(year)
        parts.append("curriculum")
        prefix = "_".join(parts)

    summary_start = find_summary_start_row(ws, header_start + 3)

    course_rows = build_course_rows(ws, summary_start, header_start + 3)
    selection_rows = build_selection_groups(course_rows)
    summary_rows = build_summary_rows(ws, summary_start)

    outdir.mkdir(parents=True, exist_ok=True)
    courses_path = outdir / f"{prefix}_courses_llm.csv"
    groups_path = outdir / f"{prefix}_selection_groups_llm.csv"
    summary_path = outdir / f"{prefix}_summary_llm.csv"
    readme_path = outdir / f"{prefix}_readme.md"

    write_csv(courses_path, COURSE_FIELDNAMES, course_rows)
    write_csv(groups_path, SELECTION_FIELDNAMES, selection_rows)
    write_csv(summary_path, SUMMARY_FIELDNAMES, summary_rows)
    write_readme(readme_path, input_path.name, sheet_name, len(course_rows), len(selection_rows), len(summary_rows))

    res_paths = {
        "courses": courses_path,
        "selection_groups": groups_path,
        "summary": summary_path,
        "readme": readme_path,
    }
    return res_paths, sheet_name, prefix


def main() -> None:
    parser = argparse.ArgumentParser(description="교육과정 편성표 xlsx를 LLM용 CSV로 변환합니다.")
    parser.add_argument("input", type=Path, help="입력 xlsx 파일 경로")
    parser.add_argument("--sheet", default=None, help="변환할 시트명 (기본값: 자동으로 탐색)")
    parser.add_argument("--outdir", type=Path, default=None, help="출력 폴더 (기본값: 입력 파일과 동일한 폴더)")
    parser.add_argument("--prefix", default=None, help="출력 파일명 prefix (기본값: 학교명과 학년도를 기준으로 자동 감지)")
    args = parser.parse_args()

    outdir = args.outdir if args.outdir else args.input.parent
    paths, sheet_name, prefix = convert(args.input, args.sheet, outdir, args.prefix)
    print(f"변환 완료 (시트: {sheet_name}, 접두사: {prefix})")
    for key, path in paths.items():
        print(f"- {key}: {path}")


if __name__ == "__main__":
    main()
